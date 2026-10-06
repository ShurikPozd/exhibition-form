"""Поиск заявок в админке и выгрузка по текущему фильтру.

Задача: организатор ищет конкретного посетителя по имени, телефону или по отметке в
анкете («МКД»), а получить хочет ровно то, что он видит на экране, — не все заявки
подряд. Поэтому проверяется и нахождение по каждому полю, и то, что фильтр живёт
в ссылках (страница, «Правка», выгрузка), и что «%» в запросе не превращается в шаблон.
"""

import io
import re
from urllib.parse import unquote

import pytest
from openpyxl import load_workbook

from services import exporters
from services import submissions as submissions_service

# Колонки выгрузки задаёт exporters.headers(), а не тест: проверка по номеру падает
# при любой перестановке заголовков, а проверка по имени — нет.
ID_COLUMN = exporters.headers().index("№")
NAME_COLUMN = exporters.headers().index("Имя")


def _names(payload: dict) -> list[str]:
    """Достаёт имена из JSON-выгрузки по колонке «Имя»."""
    return [row[NAME_COLUMN] for row in payload["rows"]]


def _create(client, payload: dict, **changes) -> int:
    """Создаёт заявку с изменениями и возвращает её номер."""
    response = client.post("/api/submissions", json={**payload, **changes})
    assert response.status_code == 201, response.text
    return int(response.json()["id"])


def _card_ids(text: str) -> list[int]:
    """Достаёт номера заявок из карточек страницы по подписи «№»."""
    parts = text.split(">№")
    found = []
    for part in parts[1:]:
        digits = ""
        for char in part:
            if char.isdigit():
                digits += char
            else:
                break
        found.append(int(digits))
    return found


@pytest.fixture()
def three_submissions(client, valid_payload):
    """Три заявки с разными контактами и разными отметками."""
    first = _create(
        client,
        valid_payload,
        name="Анна Иванова",
        company="ООО Ромашка",
        phone="+7 999 111-22-33",
        email="anna@example.com",
        directions=["МКД"],
    )
    second = _create(
        client,
        valid_payload,
        name="Пётр Сидоров",
        company='ООО "Ромашка"',
        phone="+7 999 444-55-66",
        email="petr@example.com",
        directions=["SmartHome"],
    )
    third = _create(
        client,
        valid_payload,
        name="Анна Петрова",
        company="ИП Кузнецов",
        phone="+7 999 777-88-99",
        email="kuz@example.com",
        directions=["SmartHome", "МКД"],
    )
    return first, second, third


def test_search_by_name(auth_client, three_submissions):
    """Поиск по имени находит нужную карточку, остальные не показывает."""
    listing = auth_client.get("/admin?query=Пётр")

    assert _card_ids(listing.text) == [three_submissions[1]]
    assert "Найдено по запросу «Пётр»: 1" in listing.text


def test_search_is_case_insensitive(auth_client, three_submissions):
    """Регистр не важен: человек ищет так, как видит имя в заявке."""
    assert _card_ids(auth_client.get("/admin?query=пётр").text) == [
        three_submissions[1]
    ]


def test_search_matches_partial_word(auth_client, three_submissions):
    """Достаточно части слова — телефон и имя вводят по кускам."""
    listing = auth_client.get("/admin?query=Иванов")

    assert _card_ids(listing.text) == [three_submissions[0]]


def test_search_by_phone(auth_client, three_submissions):
    """Поиск по телефону в любой части номера."""
    listing = auth_client.get("/admin?query=777-88")

    assert _card_ids(listing.text) == [three_submissions[2]]


def test_search_by_email(auth_client, three_submissions):
    """Поиск по почте находит заявку."""
    listing = auth_client.get("/admin?query=petr@example.com")

    assert _card_ids(listing.text) == [three_submissions[1]]


def test_search_by_company(auth_client, three_submissions):
    """Поиск по компании находит обе заявки этого работодателя."""
    listing = auth_client.get("/admin?query=Ромашка")

    assert _card_ids(listing.text) == [three_submissions[1], three_submissions[0]]


def test_search_by_checked_answer(auth_client, three_submissions):
    """Отметка из анкеты тоже находится словом запроса — «МКД» в payload."""
    listing = auth_client.get("/admin?query=МКД")

    assert _card_ids(listing.text) == [three_submissions[2], three_submissions[0]]


def test_search_without_match(auth_client, three_submissions):
    """Нет совпадений — пустой список с понятной подсказкой, а не «Заявок пока нет»."""
    listing = auth_client.get("/admin?query=Кого-то-другого")

    assert _card_ids(listing.text) == []
    assert "ничего не нашлось" in listing.text


def test_wildcards_are_not_patterns(auth_client, client, valid_payload):
    """«%» и «_» в запросе — обычные символы, а не шаблон LIKE."""
    _create(client, valid_payload, name="Скидка 100%", company="")
    _create(client, valid_payload, name="Обычная заявка", company="")

    assert len(_card_ids(auth_client.get("/admin?query=100%").text)) == 1
    assert len(_card_ids(auth_client.get("/admin?query=%").text)) == 1


def test_empty_query_shows_everything(auth_client, three_submissions):
    """Пустой запрос — это отсутствие фильтра, а не поиск по пустоте."""
    listing = auth_client.get("/admin?query=")

    assert len(_card_ids(listing.text)) == 3


def test_query_is_trimmed(auth_client, three_submissions):
    """Пробелы и их повторы вокруг запроса не мешают найти заявку."""
    listing = auth_client.get("/admin?query=%20%20Пётр%20%20")

    assert _card_ids(listing.text) == [three_submissions[1]]


def test_search_and_pagination_work_together(auth_client, client, valid_payload):
    """Счётчики и страницы считаются по результату поиска, а не по всей базе."""
    for number in range(15):
        _create(client, valid_payload, name=f"Тестовая {number:02d}", company="")
    for number in range(10):
        _create(client, valid_payload, name=f"Другая {number:02d}", company="")

    first_page = auth_client.get("/admin?query=Тестовая&per_page=10")
    assert len(_card_ids(first_page.text)) == 10
    assert "Показано 1–10 из 15" in first_page.text
    assert "Страница 1 из 2" in first_page.text

    second_page = auth_client.get("/admin?query=Тестовая&page=2")
    assert len(_card_ids(second_page.text)) == 5
    assert "Показано 11–15 из 15" in second_page.text


def test_search_in_hidden_list(auth_client, client, valid_payload):
    """В списке скрытых поиск работает так же и не вытаскивает видимые заявки."""
    hidden = _create(client, valid_payload, name="Скрытый посетитель")
    visible = _create(client, valid_payload, name="Скрытый посетитель два")
    auth_client.post(f"/admin/submissions/{hidden}/delete")

    listing = auth_client.get("/admin?show_deleted=true&query=Скрытый")

    assert _card_ids(listing.text) == [hidden]
    assert visible not in _card_ids(listing.text)


def test_filter_survives_pagination_links(auth_client, three_submissions):
    """Ссылки пейджера несут запрос: переход на вторую страницу не сбрасывает фильтр."""
    for number in range(12):
        _create(
            auth_client, {"name": "Пётр", "consent": True, "phone": "+7 900 000-00-00"}
        )

    listing = auth_client.get("/admin?query=Пётр")
    assert "Вперёд" in listing.text
    assert "/admin?page=2&amp;query=%D0%9F%D1%91%D1%82%D1%80" in listing.text


def test_filter_survives_edit_link(auth_client, three_submissions):
    """Кнопка «Правка» возвращает на ту же отфильтрованную страницу, а не в общий список."""
    listing = auth_client.get("/admin?query=Пётр")

    match = re.search(r'href="/admin/submissions/\d+/edit\?next=([^"]+)"', listing.text)
    assert match, "в карточке должна быть ссылка «Правка» с адресом возврата"
    target = unquote(match.group(1))
    assert target.startswith("/admin?page=1&query=")
    assert unquote(target.split("query=", 1)[1]) == "Пётр"


def test_per_page_redirect_keeps_query(auth_client, three_submissions):
    """Смена размера страницы не выбрасывает запрос из адреса."""
    response = auth_client.get("/admin?query=Пётр&per_page=50", follow_redirects=False)

    assert response.status_code == 303
    assert "query" in response.headers["location"]


def test_reset_link_visible_only_with_query(auth_client, three_submissions):
    """Кнопка сброса есть только когда есть что сбрасывать."""
    assert "Сбросить" in auth_client.get("/admin?query=Пётр").text
    assert "Сбросить" not in auth_client.get("/admin").text


def test_export_buttons_carry_current_filter(auth_client, three_submissions):
    """Кнопки выгрузки отдают текущий фильтр, а не все заявки."""
    listing = auth_client.get("/admin?query=Пётр")

    assert "/api/submissions.xlsx?query=%D0%9F%D1%91%D1%82%D1%80" in listing.text
    assert "/api/submissions.csv?query=%D0%9F%D1%91%D1%82%D1%80" in listing.text


def test_export_buttons_without_filter_stay_plain(auth_client, three_submissions):
    """Без поиска адрес выгрузки не меняется — лишних параметров в закладках не будет."""
    listing = auth_client.get("/admin")

    assert 'href="/api/submissions.xlsx"' in listing.text


def test_json_export_respects_query(auth_client, three_submissions):
    """JSON-выгрузка по токену тоже умеет фильтр."""
    response = auth_client.get("/api/submissions?query=Пётр")

    assert response.status_code == 200
    assert response.json()["count"] == 1
    assert _names(response.json()) == ["Пётр Сидоров"]


def test_xlsx_export_contains_only_matched(auth_client, three_submissions):
    """В Excel по фильтру попадают только совпавшие заявки."""
    response = auth_client.get("/api/submissions.xlsx?query=МКД")

    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content))
    sheet = workbook["Заявки"]
    names = [
        sheet.cell(row=row, column=NAME_COLUMN + 1).value
        for row in range(2, sheet.max_row + 1)
    ]
    assert sorted(names) == ["Анна Иванова", "Анна Петрова"]


def test_csv_export_respects_query(auth_client, three_submissions):
    """CSV по фильтру — это те же строки, что и в Excel."""
    response = auth_client.get("/api/submissions.csv?query=Пётр")

    assert response.status_code == 200
    text = response.content.decode("utf-8-sig")
    assert "Пётр Сидоров" in text
    assert "Анна Иванова" not in text


def test_export_can_include_hidden(auth_client, client, valid_payload):
    """Выгрузка «скрытых» отдаёт только скрытые заявки, обычная — только видимые."""
    hidden = _create(client, valid_payload, name="Скрытая заявка")
    visible = _create(client, valid_payload, name="Видимая заявка")
    auth_client.post(f"/admin/submissions/{hidden}/delete")

    plain = auth_client.get("/api/submissions?query=аявка")
    with_hidden = auth_client.get("/api/submissions?query=аявка&show_deleted=true")

    assert [row[ID_COLUMN] for row in plain.json()["rows"]] == [str(visible)]
    assert [row[ID_COLUMN] for row in with_hidden.json()["rows"]] == [str(hidden)]


def test_export_filters_do_not_break_default(client, three_submissions):
    """Без параметров выгрузка ведёт себя как раньше: все видимые заявки."""
    response = client.get("/api/submissions.xlsx?token=test-export-token")

    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook["Заявки"].max_row == 4


def test_search_normalization_rules():
    """Нормализация запроса: обрезка краёв, один пробел, ограничение длины."""
    assert submissions_service.normalize_query("  Пётр   Сидоров ") == "Пётр Сидоров"
    assert submissions_service.normalize_query("") == ""
    assert submissions_service.normalize_query(None) == ""
    assert len(submissions_service.normalize_query("а" * 500)) == 120
