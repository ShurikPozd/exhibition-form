"""Проверки выгрузки: колонки, строки, содержимое .xlsx и .csv."""

import csv
import io

from openpyxl import load_workbook

from config.form_fields import FORM_FIELDS
from schemas import SubmissionIn
from services import exporters, submissions


def make_submission(session, valid_payload, **overrides):
    """Создаёт заявку в базе — источник данных для выгрузки."""
    payload = dict(valid_payload)
    payload.update(overrides)
    return submissions.create_submission(session, SubmissionIn(**payload))


def headers_index(header: str) -> int:
    """Позиция колонки по её заголовку."""
    return exporters.headers().index(header)


def test_headers_follow_form_order(session, valid_payload):
    """Колонки идут в порядке анкеты: №, дата и время, затем поля формы."""
    headers = exporters.headers()
    assert headers[0] == "№"
    assert headers[1] == "Дата и время"
    assert len(headers) == len(FORM_FIELDS) + 2
    assert "Имя" in headers
    assert "Согласие ПДн" in headers


def test_row_contains_values(session, valid_payload):
    """В строке лежат значения всех полей, чекбоксы склеены через «;»."""
    saved = make_submission(session, valid_payload)

    row = exporters.rows([saved])[0]

    assert row[0] == str(saved.id)
    assert row[2] == "Иван Петров"
    assert "SmartHome" in row[6] and "МКД" in row[6]
    assert row[headers_index("Интересующие направления")] == "SmartHome; МКД"
    assert row[headers_index("Согласие ПДн")].startswith("да (версия")


def test_other_option_is_visible_in_export(session, valid_payload):
    """Раскрытое «другое» попадает в выгрузку отдельным пунктом."""
    saved = make_submission(
        session, valid_payload, role=["другое"], role_other="журналист"
    )

    row = exporters.rows([saved])[0]

    assert row[headers_index("Кто вы?")] == "другое: журналист"


def test_detail_rows_split_options(session, valid_payload):
    """Длинный лист разворачивает отметки в отдельные строки — удобно считать сводку."""
    saved = make_submission(session, valid_payload)

    detail = exporters.detail_rows([saved])
    questions = {row[2] for row in detail}
    options = {row[3] for row in detail}

    assert "Интересующие направления" in questions
    assert {"SmartHome", "МКД"} <= options
    assert all(len(row) == 4 for row in detail)


def test_xlsx_has_two_sheets(session, valid_payload):
    """В книге два листа: «Заявки» и «Подробно», шапка заморожена."""
    saved = make_submission(session, valid_payload)

    workbook = load_workbook(io.BytesIO(exporters.to_xlsx([saved])))

    assert workbook.sheetnames == ["Заявки", "Подробно"]
    sheet = workbook["Заявки"]
    assert sheet.freeze_panes == "A2"
    assert [cell.value for cell in sheet[1]] == exporters.headers()
    assert sheet.cell(row=2, column=3).value == "Иван Петров"


def test_csv_is_excel_friendly(session, valid_payload):
    """CSV в UTF-8 с BOM и разделителем «;» — иначе русский Excel ломает кириллицу."""
    saved = make_submission(session, valid_payload)

    content = exporters.to_csv([saved])

    assert content.startswith(b"\xef\xbb\xbf")

    decoded = content.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(decoded), delimiter=";"))
    assert rows[0] == exporters.headers()
    assert "Иван Петров" in rows[1]
    assert "ООО Ромашка" in rows[1]


def test_download_name_has_extension():
    """Имя файла выгрузки с датой и нужным расширением."""
    assert exporters.download_name("xlsx").endswith(".xlsx")
    assert exporters.download_name("csv").endswith(".csv")
    assert exporters.download_name("xlsx").startswith("zayavki-")
