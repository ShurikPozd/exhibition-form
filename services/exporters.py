"""Выгрузка заявок в Excel (.xlsx) и CSV.

Колонки строятся из описания анкеты (config/form_fields.py), поэтому порядок колонок
одинаков в Excel, CSV, JSON и Google Sheets — правится в одном месте.

CSV сохраняется в UTF-8 с BOM и разделителем «;»: в этом виде Excel на Windows открывает
кириллицу без плясок с кодировками, а русский Excel по умолчанию ждёт именно «;».
"""

import csv
import io
import logging
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config.form_fields import (
    CHECKBOX_GROUP,
    CONSENT,
    FORM_FIELDS,
    checkbox_fields,
)

logger = logging.getLogger(__name__)

JOINER = "; "
OTHER_OPTION = "другое"
CSV_DELIMITER = ";"
LIST_SHEET = "Заявки"
DETAIL_SHEET = "Подробно"
DATE_FORMAT = "%Y-%m-%d %H:%M:%S"
VALUE_MAX_LENGTH = 1000

# Поля, которые лежат отдельными колонками в таблице submissions, а не в payload
MODEL_COLUMNS = {"name", "company", "phone", "email"}

ID_HEADER = "№"
DATE_HEADER = "Дата и время"
DETAIL_HEADERS = [ID_HEADER, DATE_HEADER, "Вопрос", "Вариант"]


def headers() -> list[str]:
    """Возвращает заголовки колонок в порядке анкеты."""
    return [ID_HEADER, DATE_HEADER] + [
        field.get("export_title", field["label"]) for field in FORM_FIELDS
    ]


def selected_options(payload: dict, field: dict) -> list[str]:
    """Возвращает отмеченные варианты поля, раскрывая пункт «другое».

    Отмеченный «другое» заменяется расшифровкой («другое: журналист»), а не добавляется
    вторым пунктом — иначе в Excel один и тот же ответ попал бы в строку дважды.
    """
    options = [str(item) for item in (payload.get(field["key"]) or [])]

    other_key = field.get("other_key")
    if not other_key:
        return options

    other_value = str(payload.get(other_key) or "").strip()
    if not other_value:
        return options

    spelled = f"{OTHER_OPTION}: {other_value}"
    if OTHER_OPTION in options:
        return [spelled if item == OTHER_OPTION else item for item in options]

    return [*options, spelled]


def _group_values(payload: dict, field: dict) -> str:
    """Собирает значение группы чекбоксов в строку: «интегратор; заказчик от юр. лица»."""
    return JOINER.join(selected_options(payload, field))


def cell_value(submission, field: dict) -> str:
    """Возвращает значение одного поля заявки для выгрузки."""
    key = field["key"]

    if key in MODEL_COLUMNS:
        value = getattr(submission, key, "") or ""
    elif field["type"] == CHECKBOX_GROUP:
        value = _group_values(submission.payload or {}, field)
    elif field["type"] == CONSENT:
        value = (
            f"да (версия {submission.consent_version})"
            if submission.consent_version
            else ""
        )
    else:
        value = (submission.payload or {}).get(key, "")

    return str(value).strip()[:VALUE_MAX_LENGTH]


def rows(items: list) -> list[list[str]]:
    """Превращает список заявок в строки выгрузки."""
    result: list[list[str]] = []
    for submission in items:
        row = [
            str(submission.id),
            submission.created_at.strftime(DATE_FORMAT),
        ]
        row.extend(cell_value(submission, field) for field in FORM_FIELDS)
        result.append(row)
    return result


def detail_rows(items: list) -> list[list[str]]:
    """Длинный формат: отдельная строка на каждый отмеченный пункт.

    Нужен, чтобы в Excel можно было построить сводку «сколько интеграторов» без
    разбора строки по разделителю.
    """
    result: list[list[str]] = []
    for submission in items:
        payload = submission.payload or {}
        stamp = submission.created_at.strftime(DATE_FORMAT)
        for field in checkbox_fields():
            for option in selected_options(payload, field):
                result.append([str(submission.id), stamp, field["label"], option])
    return result


def _write_table(sheet, columns: list[str], rows_data: list[list[str]]) -> None:
    """Заполняет лист: шапка жирным на сером фоне, ниже данные."""
    sheet.append(columns)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E8E8E8")
        cell.alignment = Alignment(vertical="center")

    for row in rows_data:
        sheet.append(row)

    sheet.freeze_panes = "A2"


def _autofit(sheet, max_width: int = 60) -> None:
    """Подбирает ширину колонок по содержимому, с ограничением."""
    for index, column in enumerate(sheet.iter_cols(), start=1):
        letter = get_column_letter(index)
        longest = 0
        for cell in column:
            if cell.value is not None:
                longest = max(longest, len(str(cell.value)))
        sheet.column_dimensions[letter].width = min(max(longest + 2, 10), max_width)


def to_xlsx(items: list) -> bytes:
    """Собирает .xlsx: лист «Заявки» и лист «Подробно» в длинном формате."""
    workbook = Workbook()

    sheet = workbook.active
    sheet.title = LIST_SHEET
    _write_table(sheet, headers(), rows(items))

    detail = workbook.create_sheet(DETAIL_SHEET)
    _write_table(detail, DETAIL_HEADERS, detail_rows(items))

    _autofit(sheet)
    _autofit(detail)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def to_csv(items: list) -> bytes:
    """Собирает CSV для Excel: UTF-8 с BOM, разделитель «;», переводы строк CRLF."""
    buffer = io.StringIO()
    writer = csv.writer(
        buffer,
        delimiter=CSV_DELIMITER,
        lineterminator="\r\n",
        quoting=csv.QUOTE_MINIMAL,
    )
    writer.writerow(headers())
    writer.writerows(rows(items))
    return buffer.getvalue().encode("utf-8-sig")


def download_name(extension: str) -> str:
    """Имя файла выгрузки с датой: zayavki-20261005-1830.xlsx."""
    stamp = datetime.now().strftime("%Y%m%d-%H%M")
    return f"zayavki-{stamp}.{extension}"
